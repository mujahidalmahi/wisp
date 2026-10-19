import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = create_element('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = create_element('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = create_element(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color='444444', sz='4', val='single'):
    tblPr = table._tbl.tblPr
    tblBorders = create_element('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = create_element(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def add_field(p, text_before, field_name, text_after=''):
    if text_before:
        r_before = p.add_run(text_before)
        r_before.font.name = 'Arial'
        r_before.font.size = Pt(8.5)
        r_before.font.color.rgb = RGBColor(80, 80, 80)
    
    run = p.add_run()
    run.font.name = 'Arial'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(80, 80, 80)
    fld1 = create_element('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = create_element('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = field_name
    fld2 = create_element('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = create_element('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(fld3)

    if text_after:
        r_after = p.add_run(text_after)
        r_after.font.name = 'Arial'
        r_after.font.size = Pt(8.5)
        r_after.font.color.rgb = RGBColor(80, 80, 80)

def set_run_font(run, name='Arial', size_pt=10, bold=False, italic=False, color_rgb=(0, 0, 0)):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_p(doc, text="", bold_prefix=None, italic=False, space_after=4, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    if bold_prefix:
        r_bp = p.add_run(bold_prefix)
        set_run_font(r_bp, bold=True)
    if text:
        r_txt = p.add_run(text)
        set_run_font(r_txt, italic=italic)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bp = p.add_run(bold_prefix)
        set_run_font(r_bp, bold=True)
    if text:
        r_txt = p.add_run(text)
        set_run_font(r_txt)
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size_pt=13.5, bold=True)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size_pt=11, bold=True)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size_pt=10, bold=True, italic=True)
    return p

def add_code_block(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, 'F5F5F5')
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    set_table_borders(tbl, color='666666', sz='4', val='single')
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    lines = code_str.strip().split('\n')
    for i, line in enumerate(lines):
        r = p.add_run(line)
        set_run_font(r, name='Consolas', size_pt=8.5)
        if i < len(lines) - 1:
            p.add_run('\n')
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.space_before = Pt(0)

def add_callout(doc, text, title="FORMAL INVARIANT"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, 'FAFAFA')
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_table_borders(tbl, color='333333', sz='6', val='single')
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"[{title}] ")
    set_run_font(r_title, size_pt=9.5, bold=True)
    
    r_body = p.add_run(text)
    set_run_font(r_body, size_pt=9.5, italic=True)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.space_before = Pt(0)

def add_table_data(doc, headers, rows, col_widths=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color='444444', sz='4', val='single')
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(h)
        set_run_font(r, size_pt=9, bold=True)
        set_cell_background(hdr_cells[i], 'EAEAEA')
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=120, right=120)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    # Body rows
    for r_idx, row in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_color = 'FFFFFF' if r_idx % 2 == 0 else 'F8F8F8'
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = ""
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(val))
            set_run_font(r, size_pt=8.5)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=120, right=120)
            row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.space_before = Pt(0)
    return tbl

def build_document():
    doc = docx.Document()
    
    # Page setup: Standard 1.0 inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        
        # Header setup
        hdr = sec.header
        hp = hdr.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = hp.add_run("WISP PLATFORM ARCHITECTURE | TECHNICAL HANDOVER SPECIFICATION [CONFIDENTIAL]")
        set_run_font(r_hdr, size_pt=8, italic=True, color_rgb=(100, 100, 100))
        
        # Footer setup
        ftr = sec.footer
        fp = ftr.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_field(fp, "WISP-SPEC-BE-AI-2026-V2 | Page ", "PAGE", " of ")
        add_field(fp, "", "NUMPAGES", "")

    # Document Header / Organization
    p_org = doc.add_paragraph()
    p_org.paragraph_format.space_before = Pt(0)
    p_org.paragraph_format.space_after = Pt(2)
    r_org = p_org.add_run("WISP TELEVISION TECHNOLOGIES — CLOUD PLATFORM & DATA ENGINEERING GROUP")
    set_run_font(r_org, size_pt=9, bold=True, color_rgb=(80, 80, 80))

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Backend Platform to AI Integration Handover Specification")
    set_run_font(r_title, size_pt=20, bold=True)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Definitive Review of main Branch Post-Commit 5091878 (Issue #4, PR #25): Service Boundaries, Temporal Context Availability, Validation Frameworks, Latency Limits, and Runtime Configuration")
    set_run_font(r_sub, size_pt=10.5, italic=True, color_rgb=(50, 50, 50))

    # Metadata Table
    meta_headers = ["Attribute", "Specification / Value"]
    meta_rows = [
        ["Document Identifier", "WISP-SPEC-BE-AI-2026-V2"],
        ["Publication Date", "October 16, 2026"],
        ["Security Classification", "Strictly Confidential — Internal Engineering Handover"],
        ["Authoring Team", "Core Backend Platform & Supabase Database Engineering Team"],
        ["Target Audience", "AI Integration Engineering Team (@wisp/ai-engine), Feature Authors, UI Leads"],
        ["Target Git Commit on main", "Commit 5091878 (Merged PR #25, Resolving Issue #4 & 19 Security Blockers)"],
        ["Commit Message", "feat(server): merge intermediate backend security, sessions, cache, voting and enterprise controls into main (Issue #4) (#25)"],
        ["Status", "APPROVED & SIGNED OFF FOR AI INTEGRATION"]
    ]
    add_table_data(doc, meta_headers, meta_rows, col_widths=[2.0, 4.5])

    # Table of Contents Overview
    add_h2(doc, "Document Overview & Navigation")
    add_bullet(doc, "Section 1: Executive Summary & Core Architectural Invariants", "1. ")
    add_bullet(doc, "Section 2: Detailed Review of main Branch at Commit 5091878 (PR #25 Milestone)", "2. ")
    add_bullet(doc, "Section 3: Architecture Overview & Service Boundary Matrix", "3. ")
    add_bullet(doc, "Section 4: Context & Data Availability Specification (The Data Contract)", "4. ")
    add_bullet(doc, "Section 5: API Endpoints, Payloads & Data Transfer Objects (DTOs)", "5. ")
    add_bullet(doc, "Section 6: Validation, Anti-Spoiler Guards & Quality Control Pipeline", "6. ")
    add_bullet(doc, "Section 7: Operational Limits, Performance SLAs & Error Handling", "7. ")
    add_bullet(doc, "Section 8: LLM Provider Router, Resilience & Fallback Hierarchy", "8. ")
    add_bullet(doc, "Section 9: Configuration Management & Environment Variables Reference", "9. ")
    add_bullet(doc, "Section 10: Step-by-Step Developer Guide for the AI Integration Team", "10. ")
    add_bullet(doc, "Section 11: Sign-off, Acceptance Criteria & Engineering Checklist", "11. ")

    # Section 1
    add_h1(doc, "Section 1: Executive Summary & Core Architectural Invariants")
    add_p(doc, 
        "This specification establishes the official engineering handover from the Wisp Core Backend Platform and Database Team "
        "to the AI Integration Team responsible for the @wisp/ai-engine workspace. Wisp powers AI-Contextual Viewing for Amazon Fire TV, "
        "allowing viewers to receive synchronous, zero-spoiler scene explanations, entity details, character motivations, and interactive quizzes "
        "via a 10-foot D-pad-driven interface with sub-second response times."
    )
    add_p(doc,
        "With the landing of commit 5091878 on the main branch, the backend has transitioned from early prototyping into a hardened, enterprise-grade "
        "production architecture. Both the Backend and AI Integration teams operate under six non-negotiable architectural invariants codified in AGENTS.md:"
    )

    add_callout(doc, 
        "Every user-facing AI action must traverse the real execution path: Fire TV UI -> Cloudflare Edge / NestJS API -> "
        "@wisp/ai-engine (Temporal Context + Real AI Provider) -> 10-Foot TV Card. Simulated responses or hardcoded prompt mock circuits "
        "are strictly prohibited across all branches.",
        title="INVARIANT 1: ZERO FAKE AI / REAL EXECUTION PIPELINE"
    )

    add_callout(doc, 
        "All AI features, prompt builders, provider adapters, and pure-TypeScript anti-spoiler temporal filters reside exclusively in "
        "ai-engine/. The AI engine must build, run, and test 100% standalone without requiring the NestJS backend or Supabase database "
        "to be online.",
        title="INVARIANT 2: MODULAR AI INDEPENDENCE (@wisp/ai-engine)"
    )

    add_callout(doc, 
        "Never pass transcript lines, scene summaries, entity facts, or character twists with end_time_ms > playbackPositionMs or "
        "reveal_time_ms > playbackPositionMs into an LLM prompt. Spoiler filtering is deterministic, mathematical, and enforced before "
        "any token generation occurs.",
        title="INVARIANT 3: DETERMINISTIC TEMPORAL SPOILER SAFETY"
    )

    add_callout(doc, 
        "Never make on-screen keyboard text entry the primary user flow. Every scene context, bootstrap payload, and AI response card "
        "must provide 2 to 4 D-pad selectable QuickActionChip items formatted with concise labels (1-60 characters) and valid action types.",
        title="INVARIANT 4: ZERO-TYPING 10-FOOT FIRST USER EXPERIENCE"
    )

    add_callout(doc, 
        "Never place GROQ_API_KEY, GEMINI_API_KEY, OPENROUTER_API_KEY, or SUPABASE_SERVICE_ROLE_KEY inside frontend/. All AI provider "
        "credentials and elevated database keys remain strictly isolated within server-side execution boundaries.",
        title="INVARIANT 5: SERVER-SIDE CREDENTIAL ISOLATION"
    )

    add_callout(doc, 
        "All changes must pass npm run security:test (and security:audit) to ensure zero spoiler leaks, zero secret exposure, and strict "
        "adherence to API and prompt boundaries before landing in main.",
        title="INVARIANT 6: MANDATORY SECURITY & SPOILER VERIFICATION"
    )

    # Section 2
    add_h1(doc, "Section 2: Detailed Review of main Branch at Commit 5091878 (PR #25 Milestone)")
    add_p(doc,
        "Commit 5091878 ('feat(server): merge intermediate backend security, sessions, cache, voting and enterprise controls into main (Issue #4) (#25)') "
        "represents the definitive baseline of the main branch. This commit merged 66 modified files (+20,386 lines, -129 lines), resolving all 19 "
        "security blockers raised in Issue #4 and PR #21 (comment 6074614748). The table below details what commit 5091878 established in main:"
    )

    git_headers = ["Engineering Domain", "Components Merged in 5091878", "Technical Architecture & Invariants Established"]
    git_rows = [
        [
            "1. Authentication & Route Guards",
            "backend/src/common/guards/device-auth.guard.ts,\nbackend/src/common/auth/device-token.service.ts",
            "Attached DeviceAuthGuard to POST /wisp/query. All AI queries now strictly require a cryptographic HMAC-SHA256 device JWT. Enforces device revocation epochs and single-use enrollment challenges."
        ],
        [
            "2. Profile Device Ownership (IDOR/BOLA)",
            "backend/src/modules/ai-gateway/ai-gateway.service.ts,\nbackend/src/modules/profiles/profiles.service.ts",
            "Enforces strict device-to-profile ownership validation in processWispQuery and voteInteraction. Verifies profile.device_id === authenticatedDeviceId before invoking @wisp/ai-engine or persisting votes."
        ],
        [
            "3. Database Rate Limiting (Fail-Closed)",
            "backend/src/common/security/database-rate-limit.guard.ts,\nsupabase/migrations/0002_temporal_context_functions.sql",
            "Enforces 30 requests per 60-second sliding window via check_device_rate_limit RPC. Adopts fail-closed policy (HTTP 503 Service Unavailable) in production on RPC errors or malformed timestamps. Identity is strictly derived from verified token or trusted client IP."
        ],
        [
            "4. Temporal AI Response Cache with SHA-256",
            "backend/src/modules/ai-gateway/ai-gateway.service.ts,\nbackend/src/common/dto/response-card-validator.ts",
            "15-second bucket temporal cache incorporates SHA-256 prompt hash to prevent cross-prompt collisions: <episodeId>:<actionType>:<profileMode>:<timeBucketSec>:<promptHash>. Cache hits must pass validateCachedResponseCard (spoilerSafeUpToMs <= playbackPositionMs) before returning in <10ms."
        ],
        [
            "5. Cloudflare Edge Gateway Bypass",
            "cloudflare/src/index.ts",
            "Edge worker explicitly bypasses edge KV caching for /wisp/query and /sessions, emitting Cache-Control: private, no-cache, no-store, must-revalidate. Strictly forwards trusted CF-Connecting-IP header to prevent IP spoofing."
        ],
        [
            "6. Symmetric Card Runtime Validation",
            "backend/src/common/dto/response-card-validator.ts",
            "validateWispResponseCard shared runtime validator guarantees headline (1-120), body (1-2000), spoilerSafeUpToMs <= playbackPositionMs, 2-4 chips (1-60 chars, resolvable actionType), max 10 bullet points (300 chars max), and quiz schema."
        ],
        [
            "7. Sessions, Bookmarks & Atomic Procedures",
            "supabase/migrations/0008_atomic_session_bookmarks_and_voting.sql,\nbackend/src/modules/sessions/sessions.service.ts",
            "Filters active unexpired sessions (.gt('expires_at', nowIso)). Adds atomic PostgreSQL procedures add_session_bookmark and vote_interaction_atomic to prevent concurrent race conditions. Caps bookmark array at 100 entries."
        ],
        [
            "8. Input Sanitization & Payload Protection",
            "backend/src/common/security/input-sanitizer.ts,\nbackend/src/main.ts, backend/src/common/security/hpp.middleware.ts",
            "Prevents prototype pollution by recursively dropping __proto__, constructor, and prototype keys; enforces strict 16 KB request body ceiling via express.json({ limit: '16kb' }); eliminates duplicate query parameters via HPP middleware."
        ],
        [
            "9. Real Embedded Postgres Integration Tests",
            "backend/test/postgres-migrations.integration.test.ts,\nbackend/test/security-blockers-regression.test.ts",
            "Executes all 8 Supabase migrations in order on embedded PostgreSQL 18 (@electric-sql/pglite), verifying anti-spoiler triggers, RLS policies, RRF ranking, and 19 security blocker regression suites (1,605 lines)."
        ]
    ]
    add_table_data(doc, git_headers, git_rows, col_widths=[1.5, 2.0, 3.0])

    # Section 3
    add_h1(doc, "Section 3: Architecture Overview & Service Boundary Matrix")
    add_p(doc,
        "Following commit 5091878, the Wisp runtime architecture establishes clear boundaries between the Edge Proxy, "
        "the NestJS Gateway, the standalone @wisp/ai-engine, and the Supabase PostgreSQL database:"
    )

    arch_diagram = (
        "+-----------------------------------------------------------------------------------------+\n"
        "| [1] CLIENT TIER: Amazon Fire TV 10-Foot UI (React Native / Android TV)                 |\n"
        "|  - D-pad spatial navigation, zero-typing quick action chips (2-4 per card)              |\n"
        "|  - Instant Stage-1 Bootstrap (<50ms overlay rendering), sends Device JWT + Playhead Ms  |\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "                                           | HTTPS / WSS\n"
        "                                           v\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "| [2] EDGE TIER: Cloudflare Edge Worker (cloudflare/src/index.ts)                         |\n"
        "|  - SSL Termination, Strict CORS Preflight (OPTIONS), R2 Open-Movie Media Streaming     |\n"
        "|  - Trusted CF-Connecting-IP attribution; Upstream Proxying to NestJS Gateway            |\n"
        "|  - Cache-Control: private, no-store bypass on /wisp/query (Zero edge-caching of cards)  |\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "                                           | Upstream Fetch\n"
        "                                           v\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "| [3] GATEWAY TIER: NestJS API Gateway & Services (backend/src/)                          |\n"
        "|  - DeviceAuthGuard: HMAC-SHA256 JWT verification, device revocation epoch checks        |\n"
        "|  - Security: 16 KB payload ceiling, HPP middleware, prototype pollution sanitization    |\n"
        "|  - DatabaseRateLimitGuard: 30 req / 60s sliding window per device (Fail-closed in prod) |\n"
        "|  - Profile Device Ownership Verification (enforces IDOR / BOLA protection)              |\n"
        "|  - 15-Second Temporal AI Response Cache check (get_or_touch_ai_cache -> <10ms hit)      |\n"
        "|  - Cache Miss: Dispatches execution to @wisp/ai-engine via HybridContextSource          |\n"
        "|  - Symmetric Response Card Validation (validateWispResponseCard) before client delivery|\n"
        "+-----------------------------------------------------------------------------------------+\n"
        "                     | Context Query                     | In-Process Execution\n"
        "                     v                                   v\n"
        "+------------------------------------+  +-------------------------------------------------+\n"
        "| [5] DATABASE TIER: Supabase PGSQL  |  | [4] AI ENGINE TIER: @wisp/ai-engine             |\n"
        "|  - 16 core relational tables       |  |  - Pure TypeScript library (Zero NestJS/DB deps)|\n"
        "|  - get_spoiler_safe_context RPC    |  |  - filterContextByTimestamp anti-spoiler filter |\n"
        "|  - hybrid_spoiler_safe_rag_search  |  |  - AiFeatureRegistry: Dynamic action-type map   |\n"
        "|  - ai_response_cache (15s bucket)  |  |  - BaseAiFeatureModule: Prompt compilation     |\n"
        "|  - trg_enforce_spoiler_gate trigger|  |  - AiProviderRouter: Groq -> Gemini -> OpenR    |\n"
        "|  - Atomic procedures (0008)        |  |  - Heuristic JSON output repair & chip coercion |\n"
        "+------------------------------------+  +-------------------------------------------------+"
    )
    add_code_block(doc, arch_diagram)

    add_p(doc,
        "The Service Boundary & Responsibility Matrix defines which tier owns each operational responsibility:"
    )

    resp_headers = ["Concern / Capability", "Cloudflare Edge", "NestJS Gateway", "@wisp/ai-engine", "Supabase PostgreSQL"]
    resp_rows = [
        ["TLS Termination & CORS", "Primary (Edge)", "Secondary (Origin)", "Not Involved", "Not Involved"],
        ["Media Asset Streaming", "Primary (Cloudflare R2)", "Not Involved", "Not Involved", "Not Involved"],
        ["Client IP Attribution", "Primary (CF-Connecting-IP)", "Consumes header", "Not Involved", "Not Involved"],
        ["Device Authentication", "Pass-through header", "Primary (DeviceTokenService)", "Not Involved", "Stores devices/tokens"],
        ["Payload Size Bounding", "Pass-through stream", "Primary (16 KB ceiling)", "Not Involved", "Not Involved"],
        ["Rate Limiting Enforcement", "Pass-through", "Primary (DatabaseRateLimitGuard)", "Not Involved", "Executes RPC counter"],
        ["Tenant / IDOR Protection", "Pass-through", "Primary (Profile ownership)", "Not Involved", "RLS security policies"],
        ["Temporal Cache Checking", "Bypassed (Zero private cache)", "Primary (15s bucket hash)", "Secondary (In-mem fallback)", "Executes get_or_touch_ai_cache"],
        ["Temporal Context Retrieval", "Not Involved", "Calls ContextService", "Consumes IContextSource", "Executes get_spoiler_safe_context"],
        ["Context Anti-Spoiler Slicing", "Not Involved", "Verifies floor header", "Primary (filterContextByTimestamp)", "Filters valid_until <= playhead"],
        ["Prompt Compilation", "Not Involved", "Not Involved", "Primary (BaseAiFeatureModule)", "Not Involved"],
        ["LLM Provider Execution", "Not Involved", "Not Involved", "Primary (AiProviderRouter)", "Not Involved"],
        ["Provider Timeout / Quarantine", "Not Involved", "Not Involved", "Primary (10s/15s/300s rules)", "Not Involved"],
        ["Response Card Validation", "Not Involved", "Primary (validateWispResponseCard)", "Primary (formatResponseCard)", "Not Involved"],
        ["Autonomous Spoiler Clamping", "Not Involved", "Not Involved", "Not Involved", "Primary (trg_fn_enforce_spoiler)"],
        ["Atomic Bookmarks & Voting", "Pass-through", "Calls atomic RPCs", "Not Involved", "Executes stored procedures (0008)"]
    ]
    add_table_data(doc, resp_headers, resp_rows, col_widths=[1.5, 1.2, 1.3, 1.3, 1.2])

    # Section 4
    add_h1(doc, "Section 4: Context & Data Availability Specification (The Data Contract)")
    add_p(doc,
        "The Supabase PostgreSQL database serves as the authoritative temporal graph store. Following commit 5091878, the schema "
        "comprises 16 relational tables and 8 applied migrations, fully tested on real PostgreSQL instances."
    )

    add_h2(doc, "4.1 The 16 Relational Tables in Supabase")
    table_headers = ["Table Name", "Primary Keys / Indices", "Key Temporal Columns", "AI Context Purpose & Relevance"]
    table_rows = [
        ["programs", "id (UUID)", "created_at", "Master catalog metadata: program title, genre, synopsis, release year, content rating."],
        ["episodes", "id (UUID), program_id", "duration_ms, air_date", "Playable episode boundaries: season number, episode number, episode title, runtime duration."],
        ["scene_contexts", "id (UUID), episode_id", "start_time_ms, end_time_ms", "Scene boundaries, locations, narrative summaries, spoiler_safe_summary, and curated quickActionChips."],
        ["transcript_segments", "id (UUID), episode_id", "start_time_ms, end_time_ms", "Millisecond-precise dialogue lines: speaker identity, spoken dialogue, visual cue, token count."],
        ["entities", "id (UUID), program_id", "created_at", "Knowledge graph nodes: character, location, concept, object with baseDescription and lore notes."],
        ["entity_appearances", "id (UUID), entity_id", "reveal_time_ms, start_time_ms", "Time-coded progressive facts. progressive_fact is revealed only when reveal_time_ms <= playhead."],
        ["entity_relationships", "id (UUID), source/target", "reveal_time_ms", "Dynamic character ties (e.g. alliances, rivalries, family secrets) locked until reveal_time_ms."],
        ["episode_quizzes", "id (UUID), episode_id", "trigger_time_ms", "Post-watch interactive trivia: question, 4 option strings, correct_index (0..3), and spoiler-safe explanation."],
        ["temporal_embeddings", "id (UUID), episode_id", "valid_from_ms, valid_until_ms", "384-d / 1536-d vector chunks used for RAG search. Filtered strictly by valid_until_ms <= playhead."],
        ["conversation_sessions", "id (UUID), device_id", "expires_at, last_active_at", "Active sessions filtered by expires_at > now(). Stores recent turns (up to 10), pinned timestamps, and up to 100 bookmarks."],
        ["ai_response_cache", "id (UUID), cache_key (UNIQUE)", "time_bucket_sec, expires_at", "15-second bucket temporal cache incorporating prompt hash. Stores response_payload, provider_used, origin_latency_ms."],
        ["interaction_history", "id (UUID), device_id", "playback_position_ms", "Audit trail of viewer queries, AI cards returned, and viewer helpfulness votes (-1, 0, 1) via atomic voting."],
        ["ai_provider_metrics", "id (UUID), provider_name", "recorded_at", "Observability metrics: provider latency, status code, error type, token count, quarantine status."],
        ["device_rate_limits", "device_id (VARCHAR)", "window_reset_at", "Atomic sliding-window rate limit counters enforcing the 30 req / 60s rule with fail-closed protection."],
        ["spoiler_audit_log", "id (UUID), episode_id", "timestamp", "Autonomous log of blocked or clamped spoiler attempts triggered by database safety triggers."],
        ["profiles", "id (UUID), device_id", "created_at", "Household viewer profiles: quick_glance, deep_lore, family_friendly with spoiler strictness settings (max 10 profiles/device)."]
    ]
    add_table_data(doc, table_headers, table_rows, col_widths=[1.5, 1.4, 1.4, 2.2])

    add_h2(doc, "4.2 Context Retrieval Stored Procedures (RPCs)")
    add_p(doc,
        "The Backend interacts with the database using high-performance stored procedures:"
    )

    add_h3(doc, "1. get_spoiler_safe_context(p_episode_id, p_playback_ms, p_lookback_ms)")
    add_p(doc,
        "Returns the complete TemporalProgramContext object required by BaseAiFeatureModule.buildPrompts(). Under the hood, "
        "the stored procedure executes a single coalesced query that enforces:"
    )
    add_bullet(doc, "active_scene: Finds the scene where start_time_ms <= p_playback_ms, ordered by start_time_ms DESC LIMIT 1. Returns sceneIndex, sceneTitle, locationName, spoilerSafeSummary, and quickActionChips.")
    add_bullet(doc, "prior_scenes: Returns an array of prior scenes where start_time_ms <= p_playback_ms, sorted ascending by sceneIndex, containing only spoiler-safe summaries.")
    add_bullet(doc, "recent_transcript: Slices dialogue where end_time_ms <= p_playback_ms AND start_time_ms >= (p_playback_ms - p_lookback_ms), sorted ascending by start_time_ms.")
    add_bullet(doc, "unlocked_entities: Extracts entities having at least one appearance with reveal_time_ms <= p_playback_ms. Returns base_description, the most recent revealed progressive_fact, and whether the entity is active in the current scene.")

    add_h3(doc, "2. get_firetv_stage1_bootstrap(p_episode_id, p_device_id, p_playback_ms, p_profile_id)")
    add_p(doc,
        "Delivers a single-roundtrip bundle in <50ms when the viewer invokes Wisp on Fire TV. Returns active profile, "
        "active scene header, 2-4 QuickActionChips, verified pre-warmed cache cards, and session history in a single call."
    )

    add_h3(doc, "3. hybrid_spoiler_safe_rag_search(p_episode_id, p_playback_ms, p_query_text, p_query_embedding, p_match_count)")
    add_p(doc,
        "Executes Reciprocal Rank Fusion (RRF) over full-text tsvector search and pgvector cosine similarity. "
        "Strictly filters out any chunk with valid_until_ms > p_playback_ms to prevent semantic vector leaks."
    )

    add_h2(doc, "4.3 Dual-Mode Operation (HybridContextSource vs InMemoryCatalogContextSource)")
    add_bullet(doc, "HybridContextSource: Used by the NestJS backend. Calls the Supabase RPC get_spoiler_safe_context first; if Supabase is offline or unconfigured, it automatically falls back to InMemoryCatalogContextSource.")
    add_bullet(doc, "InMemoryCatalogContextSource: Used by standalone AI testing (engine.test.ts) and the CLI playground (playground.ts). Reads embedded fixtures from embedded-catalog.ts and executes pure-TS filtering via filterContextByTimestamp().")

    # Section 5
    add_h1(doc, "Section 5: API Endpoints, Payloads & Data Transfer Objects (DTOs)")
    add_p(doc,
        "Following commit 5091878, all endpoints strictly enforce authentication, input sanitization, and parameter bounding."
    )

    add_h2(doc, "5.1 Primary AI Query Endpoint: POST /api/v1/wisp/query")
    add_p(doc, "Invokes contextual AI feature generation or returns a cached response card.", bold_prefix="Endpoint Specification: ")
    add_bullet(doc, "Authentication: Mandatory Device Token (Authorization: Bearer <device_jwt>). Handled by DeviceAuthGuard.")
    add_bullet(doc, "Tenant Isolation: Profile ownership is verified (profile.device_id === authenticatedDeviceId). Rejects cross-device access with HTTP 403 Forbidden.")
    add_bullet(doc, "Rate Limiting: 30 requests per 60-second window per authenticated device. Fails closed (HTTP 503) in production if limiter unavailable.")
    add_bullet(doc, "Payload Size Boundary: Maximum 16 KB request body via express.json({ limit: '16kb' }).")

    add_p(doc, "WispQueryRequestDto Field Specification:", bold_prefix="Request Schema: ")
    req_headers = ["Field Name", "Type", "Validation Constraints", "Description"]
    req_rows = [
        ["episodeId", "string (UUID)", "IsNotEmpty, MaxLength(100)", "Target playable episode UUID."],
        ["playbackPositionMs", "number (int)", "Min(0), Max(86400000)", "Current media playhead position in milliseconds (0 to 24 hours)."],
        ["actionType", "string (enum/extensible)", "IsEnum(WispActionType)", "Feature module to invoke: contextual_qa, entity_explain, scene_explain, episode_recap, post_watch_quiz, or custom."],
        ["prompt", "string", "IsNotEmpty, MaxLength(300)", "Viewer prompt, chip question, or follow-up query."],
        ["profileMode", "string (enum)", "Optional, default 'quick_glance'", "Household profile mode: quick_glance (concise), deep_lore (detailed), family_friendly (G-rated)."],
        ["profileId", "string (UUID)", "Optional, MaxLength(100)", "Viewer profile UUID. Verified against authenticated device token to prevent IDOR."]
    ]
    add_table_data(doc, req_headers, req_rows, col_widths=[1.5, 1.2, 1.8, 2.0])

    add_p(doc, "WispResponseCardDto Field Specification:", bold_prefix="Response Schema: ")
    res_headers = ["Field Name", "Type", "Boundary Constraints", "Description"]
    res_rows = [
        ["headline", "string", "1 to 120 characters, non-empty", "Crisp card title optimized for 10-foot TV viewing."],
        ["body", "string", "1 to 2000 characters, non-empty", "Primary contextual narrative response."],
        ["bulletPoints", "string[]", "Optional, max 10 items, <=300 chars each", "Structured supporting takeaways."],
        ["spoilerSafeUpToMs", "number (int)", "0 <= spoilerSafeUpToMs <= playbackPositionMs", "Temporal anti-spoiler watermark. Never exceeds requested playhead."],
        ["profileMode", "string", "Matches requested profileMode", "Household tone applied to response."],
        ["providerUsed", "string", "groq | gemini | openrouter | ollama", "Name of the active LLM provider adapter that fulfilled generation."],
        ["cached", "boolean", "true on cache hit, false on origin LLM", "Indicates if payload was served from the 15-second temporal cache."],
        ["latencyMs", "number (int)", "Non-negative integer", "Total roundtrip latency in milliseconds (<10ms cached, <800ms origin)."],
        ["followUpChips", "QuickActionChip[]", "2 to 4 items, label 1..60 chars", "D-pad selectable follow-up chips ensuring zero-typing continuity."],
        ["quizPayload", "object", "Optional (mandatory for post_watch_quiz)", "Interactive quiz: question, exact 4 options, correctIndex (0..3), explanation."]
    ]
    add_table_data(doc, res_headers, res_rows, col_widths=[1.5, 1.2, 1.8, 2.0])

    add_h2(doc, "5.2 Ancillary AI Endpoints")
    add_p(doc, "1. GET /api/v1/wisp/features: Dynamic feature discovery endpoint listing all registered actions.")
    add_p(doc, "2. GET /api/v1/context/:episodeId: Instant overlay context endpoint with X-Wisp-Temporal-Floor-Ms header (<50ms).")
    add_p(doc, "3. GET /api/v1/context/:episodeId/bootstrap: Single-roundtrip Fire TV bootstrap payload.")
    add_p(doc, "4. PATCH /api/v1/wisp/interactions/:interactionId/vote: Helpful feedback voting (vote: -1 | 0 | 1) utilizing atomic stored procedure vote_interaction_atomic.")

    # Section 6
    add_h1(doc, "Section 6: Validation, Anti-Spoiler Guards & Quality Control Pipeline")
    add_p(doc,
        "Commit 5091878 finalized the six-tier defense-in-depth validation pipeline:"
    )

    val_headers = ["Validation Tier", "Component / Guard", "Inspection Logic", "Failure / Remediation Action"]
    val_rows = [
        [
            "Tier 1: Ingress Protection",
            "NestJS main.ts ValidationPipe & BodyParser",
            "Enforces 16 KB payload ceiling; strips non-whitelisted properties; rejects prototype pollution (__proto__, constructor, prototype); rejects duplicate parameters (HPP).",
            "HTTP 400 Bad Request with descriptive JSON error card."
        ],
        [
            "Tier 2: Identity & Tenant",
            "DeviceAuthGuard & Profile Ownership",
            "Cryptographically validates HMAC-SHA256 JWT; checks device revocation set; verifies profile.device_id matches authenticated token device_id.",
            "HTTP 401 Unauthorized or HTTP 403 Forbidden."
        ],
        [
            "Tier 3: Context Anti-Spoiler",
            "Supabase RPC & filterContextByTimestamp",
            "Deterministic filter: scene startTimeMs <= playhead, transcript endTimeMs <= playhead, entity revealTimeMs <= playhead. Future data is excluded before prompt compilation.",
            "Future data never enters prompt tokens (zero hallucination risk)."
        ],
        [
            "Tier 4: Model Parser Repair",
            "output-validation.ts in @wisp/ai-engine",
            "Strips markdown code fences (```json); applies heuristic repairs to broken JSON; normalizes chips; coerces invalid action types to valid registered types.",
            "Repairs in-flight or falls back to curated scene chips."
        ],
        [
            "Tier 5: Runtime Card Validation",
            "validateWispResponseCard (Backend & Engine)",
            "Strict shared validator applied symmetrically: headline (1-120), body (1-2000), spoilerSafeUpToMs <= playbackPositionMs, 2-4 chips, quiz structure.",
            "Rejects corrupt cards before client return or cache persistence."
        ],
        [
            "Tier 6: Database Trigger Gate",
            "trg_fn_enforce_interaction_spoiler_gate",
            "Autonomous PostgreSQL trigger on interaction_history. If spoiler_safe_up_to_ms > playback_position_ms, automatically clamps to playback_position_ms.",
            "Logs security audit row to spoiler_audit_log."
        ]
    ]
    add_table_data(doc, val_headers, val_rows, col_widths=[1.4, 1.5, 2.2, 1.4])

    add_callout(doc,
        "validateWispResponseCard is applied at FOUR critical operational boundaries: (1) Output generated by @wisp/ai-engine before returning "
        "to client, (2) Output generated by @wisp/ai-engine before persisting to database cache, (3) Cache hits read from Supabase before "
        "returning to client, and (4) Client-submitted bookmark cards. Any violation immediately aborts persistence.",
        title="CRITICAL QUALITY ASSURANCE GATE"
    )

    # Section 7
    add_h1(doc, "Section 7: Operational Limits, Performance SLAs & Error Handling")
    add_p(doc,
        "To satisfy the 10-foot TV viewing experience, the backend enforces rigorous latency budgets, sliding-window rate limits, "
        "and token boundaries."
    )

    add_h2(doc, "7.1 Performance Service Level Agreements (SLAs)")
    sla_headers = ["Operation / Pipeline Step", "Target P50", "Target P95", "Hard Timeout Ceiling", "Enforcing Component"]
    sla_rows = [
        ["15-Second Temporal Cache Hit", "< 5 ms", "< 10 ms", "100 ms", "Supabase get_or_touch_ai_cache / Memory"],
        ["Fire TV Stage-1 Bootstrap", "< 30 ms", "< 50 ms", "500 ms", "ContextService / get_firetv_stage1_bootstrap"],
        ["Instant Scene Context Overlay", "< 25 ms", "< 45 ms", "500 ms", "ContextController / get_spoiler_safe_context"],
        ["Origin AI Generation (Groq Llama 3.3)", "< 550 ms", "< 800 ms", "10,000 ms", "GroqModelAdapter via AiProviderRouter"],
        ["Total TV User Query Roundtrip", "< 700 ms", "< 1,100 ms", "15,000 ms", "Global Timeout (AiProviderRouter)"],
        ["Viewer Helpfulness Vote Persistence", "< 15 ms", "< 30 ms", "1,000 ms", "vote_interaction_atomic Stored Procedure"]
    ]
    add_table_data(doc, sla_headers, sla_rows, col_widths=[2.0, 1.1, 1.1, 1.1, 1.2])

    add_h2(doc, "7.2 Rate Limiting Rules & Payload Ceilings")
    add_bullet(doc, "Device / IP Sliding Window Rate Limit: Maximum 30 requests per 60-second window. Enforced by DatabaseRateLimitGuard calling check_device_rate_limit(p_device_id, 60, 30). Unauthenticated callers are bounded by IP; authenticated callers are bounded by device token.")
    add_bullet(doc, "Fail-Closed Policy in Production: If the database rate-limiting RPC fails or returns malformed timestamps, the guard fails closed and throws HTTP 503 Service Unavailable to protect the origin LLM pipeline from exhaustion.")
    add_bullet(doc, "Rate Limit Breach (HTTP 429): Returns a 10-foot compliant error card with headline 'Too Many Requests', body explaining the cooling period, and zero-typing follow-up chips.")
    add_bullet(doc, "Request Headers Emitted: Every response includes X-RateLimit-Limit: 30, X-RateLimit-Remaining: <count>, and X-RateLimit-Reset: <iso_timestamp>.")
    add_bullet(doc, "Global Payload Ceiling: Maximum 16 KB for all JSON and URL-encoded requests via express.json({ limit: '16kb' }).")
    add_bullet(doc, "Prompt Length Ceiling: Maximum 300 characters for user prompts.")
    add_bullet(doc, "Card Size Ceilings: Headline max 120 chars; Body max 2000 chars; Follow-up chips 2 to 4 items with labels 1 to 60 chars.")

    add_h2(doc, "7.3 Token Budget & Prompt Allocation")
    add_p(doc,
        "Prompts compiled by BaseAiFeatureModule operate within a strict 1,300 input token budget to guarantee sub-second "
        "inference times on primary LLM hardware:"
    )
    add_bullet(doc, "System Instructions & Rules: ~300 tokens (codifies JSON schema, profile tone, anti-spoiler constraints).")
    add_bullet(doc, "Active Scene Context: ~150 tokens (title, location, summary, current chips).")
    add_bullet(doc, "Prior Scene Summaries: ~200 tokens (chronological spoiler-safe recap).")
    add_bullet(doc, "Recent Dialogue Transcript: ~400 tokens (last 2-3 minutes of spoken dialogue).")
    add_bullet(doc, "Unlocked Entity Knowledge: ~200 tokens (revealed progressive facts for characters present).")
    add_bullet(doc, "User Prompt / Question: ~50 tokens.")
    add_bullet(doc, "Model Output Budget (maxTokens): Set strictly to 450 tokens. Enforces concise, legible response cards.")

    # Section 8
    add_h1(doc, "Section 8: LLM Provider Router, Resilience & Fallback Hierarchy")
    add_p(doc,
        "The AI Engine features a resilient, multi-provider fallback router (AiProviderRouter in ai-engine/src/providers/provider-router.ts). "
        "The router cascades sequentially through configured adapters until a valid completion is returned."
    )

    prov_headers = ["Priority", "Provider Adapter", "Target Model", "Default Timeout", "Quarantine Policy"]
    prov_rows = [
        ["1 (Primary)", "GroqModelAdapter", "llama-3.3-70b-versatile", "10,000 ms", "300s quarantine on 401/403/404"],
        ["2 (Secondary)", "GeminiModelAdapter", "gemini-2.5-flash / gemini-1.5-flash", "10,000 ms", "300s quarantine on 401/403/404"],
        ["3 (Tertiary)", "OpenRouterModelAdapter", "meta-llama/llama-3.3-70b-instruct", "10,000 ms", "300s quarantine on 401/403/404"],
        ["4 (Offline)", "OllamaModelAdapter", "llama3.2 (local http://localhost:11434)", "10,000 ms", "300s quarantine on connection faults"]
    ]
    add_table_data(doc, prov_headers, prov_rows, col_widths=[1.1, 1.5, 1.8, 1.1, 1.0])

    add_h2(doc, "8.1 Provider Timeout & Quarantine Rules")
    add_bullet(doc, "Per-Provider Attempt Budget (timeoutMs): Defaults to 10,000 ms. If a provider fails to return within 10 seconds, its AbortSignal fires, closing the HTTP socket immediately, and execution cascades to the next provider.")
    add_bullet(doc, "Global Attempt Budget (globalTimeoutMs): Defaults to 15,000 ms. Total elapsed wall-clock time across all cascading providers cannot exceed 15 seconds.")
    add_bullet(doc, "Minimum Attempt Budget (minAttemptMs): Defaults to 3,000 ms. If less than 3,000 ms remains in the global budget, the router will not attempt the next provider and throws AiGlobalTimeoutError.")
    add_bullet(doc, "Provider Quarantine (quarantineMs): Defaults to 300,000 ms (5 minutes). When a provider encounters invalid credentials (HTTP 401/403) or a missing model (HTTP 404), it is quarantined for 5 minutes. Subsequent requests bypass it immediately, saving failed network roundtrips. Quarantine expires automatically so rotated keys recover without restart.")

    # Section 9
    add_h1(doc, "Section 9: Configuration Management & Environment Variables Reference")
    add_p(doc,
        "The table below details all environment variables across backend services, edge workers, database connections, and AI providers:"
    )

    env_headers = ["Variable Name", "Service Layer", "Required?", "Default / Example", "Security Classification"]
    env_rows = [
        ["PORT", "NestJS Gateway", "Optional", "3000", "Public"],
        ["NODE_ENV", "All Services", "Mandatory", "development | production | test", "Public"],
        ["DEVICE_AUTH_SECRET", "NestJS Gateway", "Mandatory in prod", "32+ char hex secret", "Critical Secret (Server-Only)"],
        ["INTERNAL_OPERATOR_TOKEN", "NestJS Gateway", "Optional", "32+ char token", "Critical Secret (Server-Only)"],
        ["DEVICE_ENROLLMENT_SECRET", "NestJS Gateway", "Optional", "32+ char token", "Critical Secret (Server-Only)"],
        ["CORS_ALLOWED_ORIGINS", "Gateway & Edge", "Mandatory in prod", "http://localhost:3000,http://10.0.2.2:3000", "Config (No * in prod)"],
        ["TRUST_PROXY", "NestJS Gateway", "Optional", "true", "Config"],
        ["SUPABASE_URL", "Gateway & DB", "Mandatory", "https://xyzcompany.supabase.co", "Server-Side Config"],
        ["SUPABASE_SERVICE_ROLE_KEY", "NestJS Gateway", "Mandatory in prod", "eyJh... (service_role)", "Critical Secret (Never in client)"],
        ["SUPABASE_ANON_KEY", "Gateway / Client", "Optional", "eyJh... (anon public)", "Public Client Key"],
        ["GROQ_API_KEY", "AI Engine", "Mandatory for Groq", "gsk_...", "Critical Secret (Server-Only)"],
        ["GROQ_MODEL", "AI Engine", "Optional", "llama-3.3-70b-versatile", "Config"],
        ["GEMINI_API_KEY", "AI Engine", "Mandatory for Gemini", "AIzaSy...", "Critical Secret (Server-Only)"],
        ["GEMINI_MODEL", "AI Engine", "Optional", "gemini-2.5-flash", "Config"],
        ["OPENROUTER_API_KEY", "AI Engine", "Mandatory for OR", "sk-or-...", "Critical Secret (Server-Only)"],
        ["OPENROUTER_MODEL", "AI Engine", "Optional", "meta-llama/llama-3.3-70b-instruct", "Config"],
        ["OLLAMA_BASE_URL", "AI Engine", "Optional", "http://localhost:11434", "Config"],
        ["OLLAMA_MODEL", "AI Engine", "Optional", "llama3.2", "Config"],
        ["ENABLE_DEMO_FALLBACK", "Gateway & Engine", "Prohibited in prod", "false", "Config (Fails boot if true in prod)"]
    ]
    add_table_data(doc, env_headers, env_rows, col_widths=[1.8, 1.1, 1.0, 1.4, 1.2])

    add_callout(doc,
        "validateEnvironment() in backend/src/common/config/environment.validator.ts executes at process startup. "
        "It guarantees that ENABLE_DEMO_FALLBACK is never set to true in production, CORS_ALLOWED_ORIGINS never contains '*', "
        "and DEVICE_AUTH_SECRET meets the minimum 32-character entropy standard. A violation aborts process startup immediately.",
        title="PRODUCTION CONFIGURATION SAFEGUARD"
    )

    # Section 10
    add_h1(doc, "Section 10: Step-by-Step Developer Guide for the AI Integration Team")
    add_p(doc,
        "This section provides actionable guides for AI developers building new features, testing adapters, and integrating "
        "with the backend."
    )

    add_h2(doc, "10.1 Running and Testing Standalone (Zero Backend / Zero DB)")
    cmd_box = (
        "# 1. Run the interactive CLI playground with live LLM adapters\n"
        "npm run ai:playground\n\n"
        "# 2. Run the standalone AI engine test suite (<10ms execution)\n"
        "node --test ai-engine/src/engine.test.ts\n\n"
        "# 3. Run output validation and repair tests\n"
        "node --test testing/ai-integration/output-validation.test.ts\n\n"
        "# 4. Run provider timeout and quarantine regression tests\n"
        "node --test testing/ai-integration/provider-timeouts.test.ts"
    )
    add_code_block(doc, cmd_box)

    add_h2(doc, "10.2 Adding a New AI Feature Module")
    add_bullet(doc, "Step 1: Create a new file in ai-engine/src/features/my-new.feature.ts extending BaseAiFeatureModule.")
    add_bullet(doc, "Step 2: Define readonly actionType = 'my_new_action' and readonly description.")
    add_bullet(doc, "Step 3: Implement getFeatureSpecificInstructions() returning the specialized model instructions.")
    add_bullet(doc, "Step 4: Register the new module in AiFeatureRegistry (ai-engine/src/features/feature-registry.ts).")

    feature_code = (
        "import { BaseAiFeatureModule } from './base.feature.ts';\n\n"
        "export class ThemeAnalysisFeature extends BaseAiFeatureModule {\n"
        "  readonly actionType = 'theme_analysis';\n"
        "  readonly description = 'Analyzes symbolic motifs and thematic undertones in the active scene.';\n\n"
        "  protected getFeatureSpecificInstructions(): string {\n"
        "    return (\n"
        "      'Analyze visual symbolism and character subtext in the active_scene.\\n' +\n"
        "      'Constraint: Reference only revealed facts; do not foreshadow future narrative twists.'\n"
        "    );\n"
        "  }\n"
        "}"
    )
    add_code_block(doc, feature_code)

    add_p(doc,
        "Dynamic Registry Propagation: Because AiFeatureRegistry acts as the single runtime source of truth, "
        "registering your feature automatically: (1) Adds the action type to the system prompt constraint, (2) Configures "
        "the response chip validator to accept followUpChips targeting your feature, (3) Exposes the action in GET /api/v1/wisp/features, "
        "and (4) Enables the Fire TV UI to dispatch it via POST /api/v1/wisp/query without any backend code modifications."
    )

    add_h2(doc, "10.3 Pre-Commit Verification Workflow")
    add_p(doc, "Before submitting pull requests to the ai-integration or main branches, run the mandatory security verification suite:")
    add_code_block(doc, "npm run security:test\nnpm run security:audit")

    # Section 11
    add_h1(doc, "Section 11: Sign-off, Acceptance Criteria & Engineering Checklist")
    add_p(doc,
        "The checklist below outlines the formal acceptance criteria agreed upon between the Backend Platform and AI Integration teams:"
    )

    check_headers = ["Verification Item", "Criteria / Standard", "Verification Method", "Sign-Off Status"]
    check_rows = [
        ["Temporal Spoiler Safety", "Zero lines with endTimeMs > playbackPositionMs enter prompts", "filterContextByTimestamp unit test", "VERIFIED [PASS]"],
        ["Symmetric Card Validation", "validateWispResponseCard enforces 1-120 headline, 2-4 chips", "response-card-validator unit tests", "VERIFIED [PASS]"],
        ["Zero-Typing UX Guarantee", "Every response card returns 2-4 D-pad quick action chips", "BaseAiFeatureModule fallback chips", "VERIFIED [PASS]"],
        ["Provider Timeout Budgets", "10s adapter timeout, 15s global timeout, 300s quarantine", "provider-timeouts.test.ts (24/24)", "VERIFIED [PASS]"],
        ["15-Second Cache Efficiency", "Temporal cache hit with SHA-256 prompt hash in <10ms", "AiGatewayService integration tests", "VERIFIED [PASS]"],
        ["Database Rate Limiting", "30 req / 60s window per device (Fail-closed in prod)", "DatabaseRateLimitGuard tests", "VERIFIED [PASS]"],
        ["Server Credential Isolation", "Zero API keys or service role secrets in frontend/", "secret-scanner.ts verification", "VERIFIED [PASS]"],
        ["Clean Node 24 Execution", "ai-engine runs pure-TS standalone via node --test", "engine.test.ts (<10ms)", "VERIFIED [PASS]"]
    ]
    add_table_data(doc, check_headers, check_rows, col_widths=[1.5, 2.2, 1.8, 1.0])

    add_p(doc, "Engineering Leadership Sign-off Block:", bold_prefix="Formal Approval & Handover Sign-Off: ")
    add_p(doc,
        "By signing below, the Core Backend Platform Team and AI Integration Engineering Team acknowledge that the service boundaries, "
        "data schemas, validation rules, operational limits, and configurations documented in this specification represent the production "
        "standard for the Wisp platform following commit 5091878 on the main branch."
    )

    sign_headers = ["Role / Representation", "Designated Lead / Engineer", "Branch Responsibility", "Approval Date"]
    sign_rows = [
        ["Backend Platform Lead", "Mujahid Al Mahi", "server / main", "October 16, 2026"],
        ["AI Integration Lead", "AI Integration Engineering Team", "ai-integration / @wisp/ai-engine", "October 16, 2026"],
        ["Database Architecture Lead", "Supabase Data Platform Team", "database / supabase", "October 16, 2026"],
        ["QA & Security Engineering Lead", "Security & Verification Team", "testing / security", "October 16, 2026"]
    ]
    add_table_data(doc, sign_headers, sign_rows, col_widths=[1.6, 1.8, 1.8, 1.3])

    # Save strictly to Desktop
    desktop_dir = os.path.join(os.path.expanduser("~"), "Desktop")
    target_primary = os.path.join(desktop_dir, "BACKEND_TO_AI_INTEGRATION_HANDOVER_COMMIT_5091878.docx")
    doc.save(target_primary)
    print(f"Document saved successfully to {target_primary}")

    target_default = os.path.join(desktop_dir, "BACKEND_TO_AI_INTEGRATION_HANDOVER.docx")
    try:
        doc.save(target_default)
        print(f"Document also updated at {target_default}")
    except PermissionError:
        print(f"Note: {target_default} is currently open in Microsoft Word. Updated file is ready at {target_primary}")

if __name__ == "__main__":
    build_document()
